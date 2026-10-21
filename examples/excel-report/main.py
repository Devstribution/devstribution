"""Groups the rows of the first sheet by one column; counts them and optionally adds up another."""

from collections import defaultdict
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook


def run(excel_file: Path, group_by: str, sum_column: Optional[str], ctx) -> dict:
    ctx.progress(0.05, f"Opening {excel_file.name}")
    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]

    group_index = _column(header, group_by)
    sum_index = _column(header, sum_column) if sum_column else None

    counts: dict = defaultdict(int)
    totals: dict = defaultdict(float)
    total_rows = max(sheet.max_row - 1, 1)
    for n, row in enumerate(rows, start=1):
        if n % 500 == 0:
            ctx.progress(0.05 + 0.85 * n / total_rows, f"Reading row {n}")
        key = row[group_index] if group_index < len(row) else None
        key = "(empty)" if key is None else str(key)
        counts[key] += 1
        if sum_index is not None and sum_index < len(row) and isinstance(row[sum_index], (int, float)):
            totals[key] += row[sum_index]
    workbook.close()

    ctx.progress(0.95, "Writing the report")
    out_book = Workbook()
    out = out_book.active
    out.title = "Summary"
    out.append([group_by, "Rows"] + ([f"Total {sum_column}"] if sum_index is not None else []))
    for key in sorted(counts, key=lambda k: (-counts[k], k)):
        out.append([key, counts[key]] + ([totals[key]] if sum_index is not None else []))
    report = ctx.output_dir / f"{excel_file.stem}-by-{_slug(group_by)}.xlsx"
    out_book.save(report)

    rows_read = sum(counts.values())
    ctx.log(f"{rows_read} rows, {len(counts)} groups")
    return {"report": report, "summary": f"{rows_read} rows in {len(counts)} groups of {group_by}"}


def _column(header: list, name: str) -> int:
    for i, h in enumerate(header):
        if h.lower() == name.strip().lower():
            return i
    raise ValueError(f"column {name!r} not found; the sheet has: {', '.join(h for h in header if h)}")


def _slug(text: str) -> str:
    return "".join(c if c.isalnum() else "-" for c in text.lower()).strip("-") or "column"
